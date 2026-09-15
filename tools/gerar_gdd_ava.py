"""Gera o GDD no modelo do professor (Tiago Castro) preenchido com AMANDACONDA."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path.home() / "Downloads" / "GDD_AMANDACONDA_Bruno_AVA.docx"
# Também salva na pasta do projeto
OUT_PROJ = Path(__file__).resolve().parents[1] / "docs" / "GDD_AVA_modelo-professor.docx"


def set_run(run, *, bold=False, size=11, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=True, size=16)
    p.space_after = Pt(6)


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, bold=True, size=13, color=RGBColor(0x1A, 0x1A, 0x1A))
    p.space_before = Pt(14)
    p.space_after = Pt(4)


def body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=11)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    set_run(r, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run(r, bold=True, size=10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_run(r, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    h1(doc, "Game Design Document")
    body(
        doc,
        "Modelo de preenchimento — entrega individual pelo AVA até 13/08/2026. "
        "Documento preenchido conforme o jogo AMANDACONDA (projeto do semestre).",
    )

    add_table(
        doc,
        ["Campo", "Preenchimento"],
        [
            ("Nome do aluno", "_______________________________  (preencher)"),
            ("Matrícula", "_______________________________  (preencher)"),
            ("Título do jogo", "AMANDACONDA"),
            ("Data", "11/08/2026"),
            (
                "Ferramenta",
                "(   ) PyGame — jogo 2D\n"
                "( X ) Unreal Engine — jogo 3D (terceira pessoa, só Blueprint)",
            ),
        ],
    )

    # 1
    h2(doc, "1. Conceito do jogo")
    body(
        doc,
        "AMANDACONDA é um action soulslike em terceira pessoa (Unreal Engine 5) numa "
        "faculdade à noite em campus semiaberto: você explora, enfrenta inimigos, pega "
        "pendrives e decide quando encarar cada chefe — liberdade total, dificuldade alta.",
    )
    body(
        doc,
        "Estilo Elden Ring: os três chefes-base estão acessíveis cedo; ir despreparado é "
        "permitido e punido. Matá-los libera os bosses finais (Marlombólico → Amandaconda) "
        "e ainda entrega uma habilidade ligada à personalidade de cada amigo "
        "(Caramelo / Eco / Espelho) — você “rouba” a característica deles. "
        "Cada chefe ensina uma lição (espaço, ritmo, duelo, caos).",
    )
    body(
        doc,
        "Homenagem a um grupo de amigos: os quatro minichefes são versões monstruosas da "
        "personalidade de cada um; a Amandaconda é o chefe final e não é piada.",
    )

    # 2
    h2(doc, "2. Público-alvo")
    body(
        doc,
        "Jogadores de PC (Windows), 16+, que já jogaram algum soulslike ou metroidvania "
        "leve (Hollow Knight, Dead Cells, Salt and Sanctuary) ou querem um combate exigente "
        "mas curto. Sessões de 45 a 75 minutos de conteúdo (2 a 4 horas com mortes). "
        "Uso em casa, teclado+mouse ou joystick; não é mobile nem sala de aula.",
    )

    # 3
    h2(doc, "3. Gênero e referências")
    body(
        doc,
        "Gênero: action soulslike 3D (campus semiaberto, exploração, soft locks), combate "
        "exigente estilo Elden Ring / Dark Souls.",
    )
    bullet(
        doc,
        "Elden Ring — mapa conectado; soft locks; “pode ir cedo, mas é difícil”; safezones "
        "como Sites of Grace; exploração recompensa.",
    )
    bullet(
        doc,
        "Dead Cells — pendrive = pergaminho; Bytes = cells; Patches ≈ mutações.",
    )
    bullet(
        doc,
        "Dark Souls / Salt and Sanctuary — vigor, rolamento, telegrafia, culpa no jogador, "
        "atalhos, morte que ensina.",
    )

    # 4
    h2(doc, "4. Narrativa")
    body(
        doc,
        "Cenário: uma faculdade que não fechou — suspendeu. De noite o campus semiaberto "
        "(hall, bandejão em greve, lab, pátio, corredores, servidores, porta de saída) vira "
        "outro lugar: luzes de emergência, máquinas que ninguém desligou, atalhos que de dia "
        "não existem. A interface fala com você como um terminal de plantão "
        "(SESSION_STUCK / RECOVER DATA / FIND EXIT). Sem avatar nomeado: a história mora no "
        "campus e nos chefes, não num diário do herói.",
    )
    body(
        doc,
        "Por que explorar: achar a saída. O problema é a pichação ao lado da placa — "
        "“não existe saída”. Você recupera Bytes (dados), fecha “tickets” dos chefes-base "
        "(qualquer ordem; ir cedo demais = morte), rouba a habilidade ligada à personalidade "
        "de cada um, libera Marlombólico nos servidores e, enfim, a porta.",
    )
    body(
        doc,
        "Os quatro amigos são lendas de corredor — processos que não morreram, versões "
        "monstruosas da própria personalidade. A comédia está neles. O mistério está nela.",
    )
    body(
        doc,
        "O enigma da Amandaconda: ninguém sabe a idade dela (dizem que estava na arca de Noé; "
        "que viu o dilúvio e o Windows XP; que instalou a primeira rede com cabo coaxial). "
        "Lenda interna do grupo vira mito do campus: ela “bate em criança e tranca no lab” "
        "(calouro some, volta com trauma e crachá de estagiário). Fala como tiozão de "
        "informática das antigas — “já tentou desligar e ligar de novo?”, senha no Post-it, "
        "café de betume, cigarro entre apocalipses. Ela não guarda a saída pra impedir: "
        "ela é o teste da saída. Root do plantão eterno. Você só sai se ela decidir que "
        "aguentou.",
    )
    body(
        doc,
        "Cena da porta: encostada, fumando, café na mão — sua jornada inteira foi o tempo "
        "de um cigarro. Uma frase de tiozão, apoia a xícara, luta. Final: você vence, ela "
        "não morre; pega a xícara e oferece (“agora você é do plantão”). Você não a "
        "derrotou — foi aceito. Créditos com os cinco.",
    )

    # 5
    h2(doc, "5. Mecânicas principais")
    add_table(
        doc,
        ["Mecânica", "Como funciona e regras"],
        [
            (
                "Escola (início)",
                "Escolha única por run: Físico (melee / arma corpo a corpo) ou Especial "
                "(projétil / bola de vírus). Só a escola escolhida fica ativa nesta entrega.",
            ),
            (
                "Ataque leve / pesado",
                "Leve: combo corpo a corpo ou tiro. Pesado: carregável. Consomem vigor. "
                "Lock-on opcional no alvo. Números-base: leve ~12 de dano (ajustáveis na UE).",
            ),
            (
                "Rolamento / esquiva",
                "Única defesa. Invencibilidade curta (i-frames). Consome vigor. Câmera 3ª pessoa.",
            ),
            (
                "Vigor",
                "Correr, rolar e atacar gastam. Regenera parado. Sem vigor = sem spam de ação.",
            ),
            (
                "Cura (energético)",
                "Latas limitadas por tentativa; beber é lento e punível. Safezone recarrega.",
            ),
            (
                "Pendrive",
                "No chão da sala de risco. Ao pegar: pausa e fork Físico (+vida / escala melee) "
                "vs Especial (+dano ranged / menos vida). Escala: dano = base × (1+taxa)^nível.",
            ),
            (
                "Bytes / KB",
                "Moeda dropada por inimigos comuns. Gasta no Servidor de backup (Patches/Drivers). "
                "Chefes não dropam Bytes.",
            ),
            (
                "Habilidade de chefe",
                "Ao matar um chefe-base você rouba a característica dele: LuanEvil→Caramelo "
                "(cola o espaço), Renanligno→Eco (golpe atrasado), Balarrals→Espelho "
                "(contra no i-frame). Equipa uma por vez na safezone. Não é loot genérico — "
                "é a personalidade do amigo virando ferramenta.",
            ),
            (
                "Morte",
                "Na sala de risco: volta ao início da sala. No chefe: respawn na porta da arena "
                "(< 2 s), pendrives já aplicados não somem.",
            ),
            (
                "Campus aberto",
                "Exploração 3ª pessoa; safezones (“grace”); atalhos. Três arenas-base "
                "acessíveis cedo. 3 bases mortas → Marlon. Marlon morto → Amandaconda. "
                "Sem menu Mega Man de ordem obrigatória.",
            ),
        ],
    )

    # 6
    h2(doc, "6. Controles e inputs")
    body(
        doc,
        "Unreal Engine 5 — terceira pessoa com Enhanced Input. Suporte completo a "
        "teclado+mouse e joystick (Xbox / PlayStation). Mesmo moveset nos dois.",
    )
    add_table(
        doc,
        ["Ação", "Teclado + mouse", "Joystick (Xbox / PlayStation)"],
        [
            ("Mover", "W A S D", "Analógico esquerdo"),
            ("Olhar / câmera", "Mouse", "Analógico direito"),
            ("Correr (gasta vigor)", "Shift (segurar)", "L3 (click do analógico) ou LT"),
            ("Pular", "Espaço", "A / X (PlayStation)"),
            ("Ataque leve", "Botão esquerdo do mouse", "RB / R1"),
            ("Ataque pesado / carregar", "Botão direito (segurar)", "RT / R2"),
            ("Rolamento (i-frames)", "Ctrl / C", "B / ○"),
            ("Lock-on", "Botão do meio / Tab", "R3 (click do analógico direito)"),
            ("Cura (energético)", "Q", "Seta baixo / D-pad ↓"),
            ("Habilidade de chefe", "E", "X / □"),
            ("Interagir", "F", "Y / △"),
            ("Pausa", "Esc", "Menu / Options"),
        ],
    )
    body(
        doc,
        "Exemplo de leitura no jogo/HUD de ajuda: Espaço / X — pular. "
        "Protótipo atual já valida movimento Third Person + ataque leve + dummy; "
        "o mapeamento de gamepad entra junto no Enhanced Input.",
    )

    # 7
    h2(doc, "7. Personagens")
    add_table(
        doc,
        ["Nome", "Papel", "Habilidades e comportamento"],
        [
            (
                "Você (sem nome)",
                "Jogável",
                "Escola Físico ou Especial; leve/pesado; roll; vigor; cura limitada; "
                "1 habilidade de chefe equipada.",
            ),
            (
                "Chatana",
                "Inimigo comum",
                "Ataque sonoro que atrapalha percepção. Ensina priorizar alvo. Frágil sozinha.",
            ),
            (
                "Portara",
                "Inimigo comum",
                "Muro lento e resistente; não atordoa. Investida telegrafada. Melhor contornar.",
            ),
            (
                "Net",
                "Inimigo comum",
                "Antena à distância; marca o jogador (dano contínuo / revela posição).",
            ),
            (
                "LuanEvil",
                "Minichefe (bandejão)",
                "Personalidade: gente boa / tanque. Chão de caramelo; ensina espaço. "
                "Recompensa: Caramelo (poça que freia).",
            ),
            (
                "Renanligno",
                "Minichefe (lab)",
                "Personalidade: fora do tempo. Fantasmas atrasados; ensina paciência. "
                "Recompensa: Eco (golpe fantasma depois).",
            ),
            (
                "Balarrals",
                "Minichefe (pátio)",
                "Personalidade: espelho / provocação. Usa o moveset do jogador. "
                "Recompensa: Espelho (contra no roll).",
            ),
            (
                "Marlombólico",
                "Penúltimo chefe (servidores)",
                "Hackeia a UI: inverte controles, mente na vida, clones. Luta de chinelo, "
                "mão no bolso. Não entrega habilidade jogável.",
            ),
            (
                "Amandaconda",
                "Chefe final (porta)",
                "Humanoide-serpente; 2 fases; 8 ataques telegrafados; cauda como arma. "
                "Não morre no final — oferece o café.",
            ),
        ],
    )

    # 8
    h2(doc, "8. Fases e progressão da dificuldade")
    add_table(
        doc,
        ["Zona / momento", "Objetivo e dificuldade"],
        [
            (
                "Tutorial",
                "Recepção. Andar, atacar, rolar, câmera. Sem hostil.",
            ),
            (
                "Campus aberto (0 bases)",
                "Bandejão / Lab / Pátio já alcançáveis. Exploração + comuns + loot. "
                "Ir na arena “errada” cedo = muito difícil (design Elden Ring).",
            ),
            (
                "Chefes-base (qualquer ordem)",
                "LuanEvil (espaço), Renanligno (ritmo), Balarrals (espelho). Cada um dá "
                "habilidade. Soft lock: 3 bases → servidores.",
            ),
            (
                "Marlombólico",
                "Zona dos servidores. Final 1 — UI que mente, controles invertidos.",
            ),
            (
                "Amandaconda",
                "Porta de saída. Final 2 — duas fases, pico soulslike.",
            ),
        ],
    )
    body(
        doc,
        "A dificuldade não é uma trilha fixa: é o jogador escolhendo risco. Combinações de "
        "Chatana + Portara + Net nos corredores fazem o “overworld” hostil sem precisar de "
        "dezenas de tipos de inimigo.",
    )

    # 9
    h2(doc, "9. Direção de arte")
    body(
        doc,
        "3D estilizado / low-poly (não realismo AAA). Faculdades à noite: fundo escuro, "
        "uma luz colorida forte por arena, objetos banais (cadeira, pipoca, crachá, MacBook) "
        "tratados com seriedade. Uma cor por chefe (âmbar, azul frio, prata, verde neon, "
        "verde-petróleo + âmbar na Amandaconda).",
    )
    body(
        doc,
        "Pipeline de arte: Unreal + Fab/biblioteca gratuita, Mixamo (animação), Tripo/Meshy "
        "(props/objetos em FBX 1K–2K), MetaHuman só se couber no PC. Inimigos comuns são "
        "objetos (Chatana, Portara, Net) — sem rig humano caro. Amandaconda: bípede com "
        "cauda longa (física/spline simples).",
    )
    body(
        doc,
        "Referências de clima: soulslikes 3D enxutos, clareza de silhueta tipo Dead Cells, "
        "iluminação dramática; cartão de entrada de chefe estilo terminal.",
    )

    # 10
    h2(doc, "10. Áudio")
    bullet(
        doc,
        "Trilha por arena/chefe — cada luta com identidade (ex.: LuanEvil: música abafada "
        "no fone e estoura na metade da vida).",
    )
    bullet(
        doc,
        "SFX de preparação distintos por ataque (soulslike: ouvir = ler o golpe).",
    )
    bullet(
        doc,
        "Chatana: ruído que atrapalha outros sons (mecânica + áudio).",
    )
    bullet(
        doc,
        "UI terminal: beeps curtos em menu, pickup de Bytes/pendrive, hit confirm.",
    )
    bullet(
        doc,
        "Fontes: Freesound / Pixabay (livres) + ajustes; sem orquestra própria no semestre.",
    )

    # 11
    h2(doc, "11. Interface e HUD")
    body(doc, "Durante a partida:")
    bullet(doc, "Vida do jogador + vigor.")
    bullet(doc, "Bytes/KB (moeda).")
    bullet(doc, "Escola ativa e níveis Físico/Especial.")
    bullet(doc, "Latas de cura restantes.")
    bullet(doc, "Habilidade de chefe equipada (ícone/texto).")
    bullet(doc, "Barra de vida do chefe + cartão de entrada na abertura da luta.")
    body(doc, "Fora da partida:")
    bullet(doc, "Menu principal, Configurações, abertura/intro rolável.")
    bullet(doc, "Pausa (Esc).")
    bullet(doc, "Escolha de escola; fork do pendrive; Servidor de backup (Patches/Drivers).")
    bullet(doc, "Safezone (grace); banco de habilidades de chefe; retry rápido na arena.")

    # 12
    h2(doc, "12. Escopo")
    add_table(
        doc,
        ["O jogo terá", "O jogo não terá"],
        [
            (
                "1 personagem jogável (você) em 3ª pessoa\n"
                "Campus semiaberto com soft locks (não open world infinito)\n"
                "2 escolas (Físico / Especial)\n"
                "3 inimigos comuns (objetos)\n"
                "3 chefes-base acessíveis + 2 finais (Marlon, Amanda)\n"
                "Safezones, atalhos, pendrives, Bytes, Patches básicos\n"
                "Combate bem soulslike (roll, vigor, telegrafia, lock-on)\n"
                "Unreal Engine 5, só Blueprint\n"
                "Assets biblioteca / Tripo / Mixamo",
                "Menu Mega Man de “escolha a ordem das portas”\n"
                "Versão 2D / PyGame nesta entrega\n"
                "Open world do tamanho de Elden Ring\n"
                "Roguelike infinito\n"
                "Multijogador / online\n"
                "Inventário / build complexa tipo RPG completo\n"
                "Escudo / aparo obrigatório\n"
                "Dublagem / cutscenes longas\n"
                "Mobile / C++\n"
                "Lumen/Nanite obrigatórios",
            ),
        ],
    )

    # 13
    h2(doc, "13. Protótipo")
    body(
        doc,
        "Pergunta: “o combate 3ª pessoa + a liberdade de ir pra zona difícil cedo funciona?”",
    )
    body(
        doc,
        "Protótipo Unreal: Third Person + ataque + dummy → rolamento/vigor/lock-on → um "
        "braço do campus (corredor + safezone + Balarrals) → segundo braço aberto sem tranca "
        "(pra provar o feel Elden Ring) → soft lock dos finais.",
    )
    body(
        doc,
        "Se o prazo apertar: 1–2 braços do campus + Balarrals + Amandaconda; os outros "
        "chefes-base com arena e moveset curto, mas já no mapa.",
    )

    # 14
    h2(doc, "14. Riscos")
    add_table(
        doc,
        ["Risco", "Mitigação"],
        [
            (
                "Escopo grande demais (5 chefes + sistemas)",
                "Priorizar Amandaconda + Balarrals bem feitos; minichefes com 3 ataques/1 fase. "
                "Perguntar ao professor se a Fase 2 pode ser fatia.",
            ),
            (
                "Combate sem “feel”",
                "Travar números em NUMEROS.md / balance.json; não avançar chefes antes do "
                "marco “bater no boneco é gostoso”.",
            ),
            (
                "Arte/animação 3D cara",
                "Inimigos = objetos (sem rig humano); Mixamo/Tripo/biblioteca; Amanda bípede "
                "com cauda simples; FBX com textura 1K–2K.",
            ),
            (
                "PC fraco pro Unreal (8 GB / Iris Xe)",
                "Projeto leve: sem Lumen/Nanite/RT; escalabilidade Low; path D:\\AmandacondaUE; "
                "cenas pequenas; fechar Chrome ao editar.",
            ),
            (
                "Perda de arquivos / OneDrive",
                "Git + Git LFS; não salvar actors Python no mapa; projeto via junction sem espaço.",
            ),
        ],
    )

    body(
        doc,
        "Antes de enviar: preencher Nome e Matrícula no topo. Relido como se outra pessoa "
        "fosse construir o jogo só com este documento.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    OUT_PROJ.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PROJ)
    print("Salvo:", OUT)
    print("Salvo:", OUT_PROJ)


if __name__ == "__main__":
    main()
